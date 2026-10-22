import io

import pytest
from docx import Document
from openpyxl import Workbook

from app import parsing, pipeline


def minimal_pdf(text: str) -> bytes:
    stream = b"BT /F1 12 Tf 72 720 Td (" + text.encode() + b") Tj ET"
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + o + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1))
    for o in offsets:
        out.write(b"%010d 00000 n \n" % o)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref))
    return out.getvalue()


def test_pdf():
    assert parsing.extract_text("brd.pdf", minimal_pdf("The system shall ship orders.")) == "The system shall ship orders."
    with pytest.raises(parsing.ParseError, match="damaged"):  # was the generic "Could not read file"
        parsing.extract_text("broken.pdf", b"%PDF-1.4 garbage")


def test_docx_headings_and_tables():
    doc = Document()
    doc.add_heading("Order capture", level=1)
    doc.add_paragraph("The system shall allow buyers to place orders.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "ID", "Requirement"
    table.cell(1, 0).text, table.cell(1, 1).text = "BR-7", "Invoices must show contract prices."
    buf = io.BytesIO()
    doc.save(buf)
    text = parsing.extract_text("brd.docx", buf.getvalue())
    assert "# Order capture" in text and "place orders" in text and "BR-7 | Invoices must show contract prices." in text
    assert pipeline.chunk_sections(text, 5000)[0]["heading"] == "Order capture"


def _xlsx(rows, title="Order capture"):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_requirement_sheet_becomes_requirements():
    data = _xlsx([["Req ID", "Requirement", "Priority"],
                  ["BR-1", "The system shall allow buyers to cancel open orders", "High"],
                  [None, None, None],
                  ["BR-2", "The system must email confirmations.", None]])
    text = parsing.extract_text("reqs.xlsx", data)
    assert text.splitlines() == ["# Order capture", "BR-1: The system shall allow buyers to cancel open orders. (Priority: High)",
                                 "BR-2: The system must email confirmations."]
    reqs = pipeline.mock_requirements(pipeline.chunk_sections(text, 5000)[0])
    assert [r["source_id"] for r in reqs] == ["BR-1", "BR-2"]


def test_xlsx_generic_table_and_csv():
    text = parsing.extract_text("t.xlsx", _xlsx([["Field", "Rule"], ["Lot", "Required for biologics"]], title="Data"))
    assert text == "# Data\nField: Lot; Rule: Required for biologics"
    csv_text = parsing.extract_text("r.csv", "ID;Description\nNFR-1;Search must respond within 2 seconds\n".encode("utf-8-sig"))
    assert csv_text == "# Requirements\nNFR-1: Search must respond within 2 seconds."


# ---- hostile and oversized files --------------------------------------------

def _zip_with_big_member(size_mb: int) -> bytes:
    """A small ZIP whose single member expands to size_mb (a decompression bomb in miniature)."""
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        with zf.open("word/document.xml", "w", force_zip64=True) as member:
            chunk = b"0" * (1024 * 1024)
            for _ in range(size_mb):
                member.write(chunk)
    return buf.getvalue()


@pytest.mark.parametrize("name", ["bomb.docx", "bomb.xlsx"])
def test_zip_bombs_are_rejected_before_anything_is_expanded(monkeypatch, name):
    from app import config
    monkeypatch.setattr(config, "MAX_UNZIPPED_MB", 5)
    data = _zip_with_big_member(20)
    assert len(data) < 200_000  # tiny on the wire
    with pytest.raises(parsing.ParseError, match="expands to"):
        parsing.extract_text(name, data)


def test_zip_with_too_many_entries_is_rejected(monkeypatch):
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for i in range(10_001):
            zf.writestr("f%d" % i, b"")
    with pytest.raises(parsing.ParseError, match="too many"):
        parsing.extract_text("many.docx", buf.getvalue())


def _blank_pdf(pages: int, password: str = "") -> bytes:
    from pypdf import PdfWriter
    w = PdfWriter()
    for _ in range(pages):
        w.add_blank_page(width=200, height=200)
    if password:
        w.encrypt(password)
    out = io.BytesIO()
    w.write(out)
    return out.getvalue()


def test_pdf_page_limit_encryption_and_scans_get_clear_messages(monkeypatch):
    from app import config
    monkeypatch.setattr(config, "MAX_PDF_PAGES", 3)
    with pytest.raises(parsing.ParseError, match="4 pages.*limit of 3"):
        parsing.extract_text("big.pdf", _blank_pdf(4))
    with pytest.raises(parsing.ParseError, match="password"):
        parsing.extract_text("locked.pdf", _blank_pdf(1, "secret"))
    with pytest.raises(parsing.ParseError, match="scanned"):
        parsing.extract_text("scan.pdf", _blank_pdf(2))
    with pytest.raises(parsing.ParseError, match="damaged"):
        parsing.extract_text("junk.pdf", b"%PDF-1.4 this is not really a pdf")


def test_damaged_office_files_say_so():
    with pytest.raises(parsing.ParseError, match="damaged"):
        parsing.extract_text("broken.docx", b"not a zip at all")
    with pytest.raises(parsing.ParseError, match="damaged"):
        parsing.extract_text("broken.xlsx", b"PK\x03\x04 truncated")


def test_truncation_is_announced_not_silent(monkeypatch):
    from app import config
    # a long document
    monkeypatch.setattr(config, "MAX_DOC_CHARS", 200)
    text = parsing.extract_text("long.txt", ("The system shall do thing number %d.\n" % 1000).encode() * 50)
    assert "truncated" in text.lower() and "200" in text
    # a long spreadsheet
    monkeypatch.setattr(config, "MAX_DOC_CHARS", 2_000_000)
    monkeypatch.setattr(parsing, "MAX_ROWS", 3)
    wb = Workbook()
    ws = wb.active
    ws.append(["ID", "Requirement"])
    for i in range(10):
        ws.append(["R%d" % i, "The system shall do thing %d." % i])
    buf = io.BytesIO()
    wb.save(buf)
    sheet = parsing.extract_text("big.xlsx", buf.getvalue())
    assert "only the first 3 rows" in sheet.lower() and "R2" in sheet and "R5" not in sheet
    # a short document is untouched
    assert "truncated" not in parsing.extract_text("ok.txt", b"The system shall export invoices.").lower()


def test_a_large_word_document_is_read_in_seconds_not_minutes():
    import time
    doc = Document()
    doc.add_heading("Order capture", level=1)
    for i in range(20000):
        doc.add_paragraph("The system shall support capability number %d for the buyer." % i)
    doc.add_heading("Invoicing", level=1)
    doc.add_paragraph("The system shall invoice at month end.")
    buf = io.BytesIO()
    doc.save(buf)
    t = time.time()
    text = parsing.extract_text("big.docx", buf.getvalue())
    took = time.time() - t
    assert took < 8, "20,000 paragraphs took %.1f s (python-docx looked the style up from scratch for every paragraph: about 25 s)" % took
    assert text.count("\n# ") >= 1 and "# Order capture" in text, "headings are still recognised"


def test_reading_a_word_document_stops_at_the_text_limit(monkeypatch):
    from app import config
    monkeypatch.setattr(config, "MAX_DOC_CHARS", 5000)
    doc = Document()
    for i in range(5000):
        doc.add_paragraph("The system shall support capability number %d for the buyer." % i)
    buf = io.BytesIO()
    doc.save(buf)
    text = parsing.extract_text("long.docx", buf.getvalue())
    assert "truncated" in text.lower() and len(text) < 20000, "no need to extract what would be cut anyway (%d chars)" % len(text)
