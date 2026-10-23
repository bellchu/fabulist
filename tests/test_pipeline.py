import json

import pytest

from app import exporters, llm, parsing, pipeline, quality

BRD = """# Order Management

## Order capture
BR-1: The system shall allow a sales representative to create orders for pharmacy customers.
The platform must track lot and expiry for biologic vaccines held in cold chain storage.
The system must respond to order searches within 2 seconds.

## Billing
Finance should be able to process chargeback claims from wholesalers quickly.
"""
OPTS = {"chunk_chars": 5000, "max_chunks_per_run": 50, "rag_top_k": 0, "redact_pii": True}
MOCK = {"kind": "mock", "name": "mock", "model": ""}


def _reqs():
    reqs = pipeline.extract_requirements([{"id": "d1", "filename": "brd.md", "text": BRD}], "sap", MOCK, OPTS, lambda m: None, {"in": 0, "out": 0})
    for i, r in enumerate(reqs):
        r["key"] = "REQ-%03d" % (i + 1)
        r["issues"] = quality.requirement_issues(r)
    return reqs


def test_chunk_sections_tracks_epic_and_feature():
    chunks = pipeline.chunk_sections(BRD, 5000)
    assert [(c["epic"], c["heading"]) for c in chunks] == [("Order Management", "Order capture"), ("Order Management", "Billing")]
    long = "# H\n" + "\n".join("line %d requirement text here" % i for i in range(400))
    assert all(len(p["text"]) <= 1000 for p in pipeline.chunk_sections(long, 1000))


def test_extract_requirements_types_ids_and_scope():
    reqs = _reqs()
    assert len(reqs) == 4
    assert reqs[0]["source_id"] == "BR-1" and reqs[0]["statement"].startswith("The system shall allow")
    assert reqs[0]["actors"] == ["sales representative"]
    assert reqs[2]["type"] == "non-functional"
    assert "Biologics and vaccines" in reqs[1]["product_scope"]
    assert reqs[1]["section"] == "Order capture" and reqs[1]["epic"] == "Order Management"


def test_extract_respects_section_limit():
    with pytest.raises(ValueError):
        pipeline.extract_requirements([{"id": "d", "filename": "f", "text": BRD}], "sap", MOCK, dict(OPTS, max_chunks_per_run=1), lambda m: None, {"in": 0, "out": 0})


def test_mock_stories_link_requirements_and_fold_nfrs():
    reqs = _reqs()
    stories = pipeline.generate_stories(reqs, "p", "sap", MOCK, OPTS, lambda m: None, {"in": 0, "out": 0})
    linked = {k for s in stories for k in s["requirement_ids"]}
    assert linked == {r["key"] for r in reqs}, "every requirement must be covered"
    nfr = next(r for r in reqs if r["type"] == "non-functional")
    host = next(s for s in stories if nfr["key"] in s["requirement_ids"])
    assert any("2 seconds" in a for a in host["acceptance_criteria"])
    cold = next(s for s in stories if "Biologics and vaccines" in s["product_scope"])
    assert any("Temperature excursions" in v["criteria"] for v in cold["variant_criteria"])
    billing = next(s for s in stories if s["feature"] == "Billing")
    assert billing["open_questions"], "vague 'quickly' should surface as an open question"


def test_requirement_quality_flags():
    issues = {i["code"] for i in quality.requirement_issues({"statement": "The system should be fast and user-friendly, TBD.", "type": "non-functional"})}
    assert {"incomplete", "ambiguous", "weak-modal", "not-verifiable"} <= issues
    assert quality.requirement_issues({"statement": "The system shall export invoices to CSV.", "type": "functional"}) == []


def test_story_quality_and_readiness():
    good = {"title": "Export invoices", "persona": "finance analyst", "narrative": "As a finance analyst, I want to export invoices, so that I can reconcile.",
            "acceptance_criteria": ["Given invoices exist, when I export, then a CSV downloads."], "requirement_ids": ["REQ-001"], "open_questions": []}
    issues = quality.story_issues(good, {"REQ-001"})
    assert issues == [] and quality.is_ready(issues) and quality.score(issues) == 100
    bad = quality.story_issues({"title": "x", "narrative": "do it", "acceptance_criteria": [], "requirement_ids": ["REQ-404"]}, {"REQ-001"})
    codes = {i["code"] for i in bad}
    assert {"narrative", "testable", "trace", "valuable"} <= codes and not quality.is_ready(bad)


def test_normalize_story_coerces_fields():
    assert pipeline.normalize_story({"title": ""}) is None
    s = pipeline.normalize_story({"title": "T", "priority": "urgent", "story_points": "5", "labels": "a"})
    assert s["priority"] == "Medium" and s["story_points"] == 5 and s["labels"] == ["a"]
    assert pipeline.normalize_story({"title": "T", "story_points": 4})["story_points"] is None


def test_parse_json_array_handles_fences_and_prose():
    assert pipeline.parse_json_array('Here:\n```json\n[{"title": "a"}]\n```') == [{"title": "a"}]
    with pytest.raises(ValueError):
        pipeline.parse_json_array("no json here")


def test_redact_masks_identifiers():
    out = llm.redact("Mail a.b@example.com or call (555) 123-4567, SSN 123-45-6789, card 4111111111111111")
    assert "@" not in out and "123-4567" not in out and "6789" not in out and "4111" not in out


def test_parsing_rejects_unsupported_and_empty():
    with pytest.raises(parsing.ParseError):
        parsing.extract_text("x.exe", b"abc")
    with pytest.raises(parsing.ParseError):
        parsing.extract_text("x.txt", b"   ")


def _story(**kw):
    base = {"id": "1", "key": "US-001", "epic": "E", "feature": "F", "title": "=cmd()", "persona": "p", "narrative": "n",
            "acceptance_criteria": ["a"], "priority": "High", "story_points": 3, "labels": [], "product_scope": [],
            "variant_criteria": [{"scope": "s", "criteria": "c"}], "compliance_tags": ["DEA"], "requirement_ids": ["REQ-001"],
            "open_questions": [], "status": "approved", "quality_score": 90, "ready": True}
    base.update(kw)
    return base


REQ = {"key": "REQ-001", "type": "functional", "statement": "The system shall x", "source_ref": "d > h", "status": "approved"}


@pytest.mark.parametrize("fmt", list(exporters.EXPORTERS))
def test_every_exporter_produces_output(fmt):
    data, mime = exporters.EXPORTERS[fmt][0]([_story()], [REQ, dict(REQ, key="REQ-002")])
    assert data and mime


def test_csv_neutralizes_formula_injection_and_rtm_flags_gaps():
    data, _ = exporters.stories_csv([_story()], [REQ])
    assert b"'=cmd()" in data and b",=cmd()" not in data
    rtm, _ = exporters.rtm_csv([_story()], [REQ, dict(REQ, key="REQ-002")])
    assert b"REQ-002" in rtm and b"GAP" in rtm
    assert json.loads(exporters.as_json([_story()], [REQ])[0])["stories"][0]["key"] == "US-001"


def test_classify_and_extract_edge_cases():
    assert pipeline.classify("The system must respond within 2 seconds for 95% of requests.") == "non-functional"
    assert pipeline.classify("Customers can request returns within 30 days of delivery.") == "rule"
    reqs = pipeline.mock_requirements({"text": "BR-303: Returned implants require UDI capture before credit is issued, rules TBD."})
    assert len(reqs) == 1 and reqs[0]["source_id"] == "BR-303"
    issues = {i["code"] for i in quality.requirement_issues({"statement": reqs[0]["statement"], "type": "rule"})}
    assert "incomplete" in issues
    assert "not-singular" in {i["code"] for i in quality.requirement_issues({"statement": "Agents should place orders and screens must be clear.", "type": "functional"})}


def test_duplicate_requirements_detected():
    reqs = [{"key": "REQ-001", "statement": "The system shall allow buyers to cancel open orders before shipment."},
            {"key": "REQ-002", "statement": "Buyers must be able to cancel open orders before shipment."},
            {"key": "REQ-003", "statement": "The system shall export invoices to CSV."},
            {"key": "REQ-004", "statement": "Short."}]
    d = quality.duplicate_pairs(reqs)
    assert d == {"REQ-001": ["REQ-002"], "REQ-002": ["REQ-001"]}


def test_html_report_escapes_and_flags_gaps():
    data, mime = exporters.html_report([_story(title="<script>x</script>")], [REQ, dict(REQ, key="REQ-002")], "Demo <b>")
    text = data.decode()
    assert mime == "text/html" and "<script>x" not in text and "&lt;script&gt;" in text and "Demo &lt;b&gt;" in text
    assert "<span class='gap'>Gap</span>" in text and "50%" in text


def test_mock_refine_requirement_splits_and_strengthens():
    req = {"key": "REQ-004", "type": "functional", "section": "Order capture", "epic": "Orders", "source_ref": "brd.md > Order capture",
           "source_id": "BR-104", "actors": ["customer service agent"],
           "statement": "Customer service agents should be able to place orders on behalf of a clinic and the order screens must be user-friendly."}
    out = pipeline.refine_requirement(req, quality.requirement_issues(req), "generic", {"kind": "mock"}, {}, {"in": 0, "out": 0})
    assert [r["statement"] for r in out] == ["Customer service agents shall be able to place orders on behalf of a clinic.",
                                             "The order screens must be user-friendly."]
    assert all(r["source_id"] == "BR-104" and r["section"] == "Order capture" for r in out)
    assert "user-friendly" in out[1]["notes"] and out[0]["notes"] == ""


def test_quality_false_positives_from_real_run():
    identify = {"key": "REQ-010", "statement": "The system shall enable Compliance to identify affected lots within 4 hours of a recall notice."}
    quarantine = {"key": "REQ-011", "statement": "The system shall enable Compliance to quarantine affected lots within 4 hours of a recall notice."}
    assert quality.duplicate_pairs([identify, quarantine]) == {}, "one different action verb = distinct requirements"
    audit = {"statement": "The system shall capture all pricing changes in an audit trail recording the user, timestamp, and reason.", "type": "non-functional"}
    assert quality.requirement_issues(audit) == []
    slow = {"statement": "Inventory availability checks must respond quickly.", "type": "non-functional"}
    assert "not-verifiable" in {i["code"] for i in quality.requirement_issues(slow)}
    noted = {"statement": "The system shall require UDI capture for returned implants.", "type": "rule", "notes": "Credit rules TBD"}
    assert "incomplete" in {i["code"] for i in quality.requirement_issues(noted)}
    assert "TBD" in pipeline.REQ_SYSTEM


def test_gherkin_steps_and_feature_zip():
    import io
    import zipfile
    assert exporters.gherkin_steps("Given I am signed in as a buyer, when I submit an order, then I see a confirmation.") == [
        "Given I am signed in as a buyer", "When I submit an order", "Then I see a confirmation"]
    assert exporters.gherkin_steps("Order status becomes Cancelled") == ["Then order status becomes Cancelled"]
    story = _story(title="Cancel orders", acceptance_criteria=["Given an open order, when I cancel it, then it is cancelled.", "Audit entry is written"],
                   variant_criteria=[{"scope": "Controlled substances", "criteria": "Cancellation is reported to DEA."}])
    data, mime = exporters.gherkin_zip([story], [REQ])
    z = zipfile.ZipFile(io.BytesIO(data))
    assert mime == "application/zip" and z.namelist() == ["features/e-f.feature"]
    text = z.read("features/e-f.feature").decode()
    assert "Feature: E / F" in text and "@US-001 @REQ-001 @DEA" in text and "    When I cancel it" in text
    assert "Then audit entry is written" in text and 'Given the product is in class "Controlled substances"' in text


def test_confluence_storage_escapes_content():
    body = exporters.confluence_storage([_story(title="<b>x</b> & y")], [REQ, dict(REQ, key="REQ-002")], "Demo")
    assert "&lt;b&gt;x&lt;/b&gt; &amp; y" in body and "<strong>Gap</strong>" in body and "50%" in body


def test_provider_failure_names_the_section_and_says_nothing_was_saved(monkeypatch):
    from app import llm
    doc = {"filename": "brd.md", "text": "# A\nThe system shall export orders.\n\n# B\nThe system shall email invoices.\n\n# C\nThe system shall archive invoices."}
    calls = []

    def flaky(profile, system, user, opts):
        calls.append(1)
        if len(calls) == 2:
            raise llm.LLMError("provider returned HTTP 503")
        return '[{"statement": "The system shall export orders to CSV.", "type": "functional"}]', {"in": 1, "out": 1}
    monkeypatch.setattr(pipeline.llm, "complete", flaky)
    opts = {"chunk_chars": 40, "max_chunks_per_run": 50, "redact_pii": False}
    with pytest.raises(llm.LLMError) as err:
        pipeline.extract_requirements([doc], "generic", {"kind": "openai"}, opts, lambda m: None, {"in": 0, "out": 0})
    msg = str(err.value)
    assert "section 2 of 3" in msg.lower() and "HTTP 503" in msg and "nothing was saved" in msg.lower()


def test_html_report_is_a_circulable_document_with_context_status_print_layout_and_sign_off():
    stories = [_story(key="US-001", status="approved"), _story(key="US-002", status="approved"), _story(key="US-003", status="draft")]
    data, _ = exporters.html_report(stories, [REQ], "Demo", details=["Generated 2026-10-21 02:41 UTC", "Scope: approved only <b>"])
    text = data.decode()
    # who/when/what: shown under the title, escaped
    assert "Generated 2026-10-21 02:41 UTC" in text and "Scope: approved only &lt;b&gt;" in text and "<b>" not in text.split("Scope:")[1][:30]
    # a status breakdown, so a reader knows how much of the backlog is agreed
    assert "2 approved" in text and "1 draft" in text
    # print: page margins and numbers, repeating table headers, rows and headings that stay together
    assert "@page" in text and "counter(page)" in text
    assert "<thead>" in text and "table-header-group" in text and "break-inside:avoid" in text.replace(" ", "")
    # sign-off block for the document that goes to the business
    assert "Review and approval" in text and text.count("Signature") >= 1 and "Date" in text


def test_html_report_without_details_still_renders():
    data, _ = exporters.html_report([_story()], [REQ], "Demo")
    assert b"Generated by Fabulist" in data and b"Review and approval" in data


def _brute_force_duplicates(reqs, threshold=0.85):
    toks = {r["key"]: quality._tokens(r["statement"]) for r in reqs}
    out = {}
    keys = list(toks)
    for i, a in enumerate(keys):
        for b in keys[i + 1:]:
            ta, tb = toks[a], toks[b]
            if len(ta) >= 3 and len(tb) >= 3 and len(ta & tb) / len(ta | tb) >= threshold:
                out.setdefault(a, []).append(b)
                out.setdefault(b, []).append(a)
    return {k: sorted(v) for k, v in out.items()}


def test_duplicate_detection_is_exact_not_approximate():
    import random
    random.seed(11)
    vocab = ["order", "invoice", "return", "lot", "expiry", "recall", "serial", "udi", "cold", "chain", "credit", "approve", "report", "audit",
             "export", "import", "sync", "account", "product", "carrier", "label", "batch", "quote", "tax", "refund", "notify", "portal", "search"]
    reqs = []
    for i in range(400):
        base = random.sample(vocab, random.randint(5, 9))
        reqs.append({"key": "REQ-%03d" % i, "statement": "The system shall " + " ".join(base) + "."})
        if i % 7 == 0:  # plant a near-duplicate: one word swapped or added
            twin = list(base)
            twin[0] = random.choice(vocab) if random.random() < .5 else twin[0]
            reqs.append({"key": "REQ-%03d-b" % i, "statement": "The system shall " + " ".join(twin + ([random.choice(vocab)] if random.random() < .3 else [])) + "."})
    expected = _brute_force_duplicates(reqs)
    assert expected, "the fixture must contain duplicates"
    assert quality.duplicate_pairs(reqs) == expected


def test_duplicate_detection_scales_to_thousands_of_requirements():
    import random
    import time
    random.seed(5)
    vocab = ["w%d" % i for i in range(3000)]
    weights = [1 / (i + 1) for i in range(3000)]
    reqs = [{"key": "REQ-%04d" % i, "statement": "The system shall " + " ".join(random.choices(vocab, weights=weights, k=random.randint(8, 14))) + "."} for i in range(4000)]
    started = time.monotonic()
    quality.duplicate_pairs(reqs)
    assert time.monotonic() - started < 0.5, "4000 requirements took %.2fs (about 1.1 s before the prefix filter)" % (time.monotonic() - started)


def test_story_batches_group_by_feature_area_and_split_at_25():
    reqs = [{"key": "REQ-%03d" % i, "type": "functional", "statement": "The system shall do thing %d." % i, "epic": "E1", "section": "A"} for i in range(30)]
    reqs += [{"key": "REQ-100", "type": "rule", "statement": "The system shall obey rule.", "epic": "E1", "section": "B"}]
    batches = pipeline.story_batches(reqs)
    assert [(e, s, len(items)) for e, s, items in batches] == [("E1", "A", 25), ("E1", "A", 5), ("E1", "B", 1)]


def test_every_provider_request_is_counted_for_the_run(monkeypatch):
    calls = []

    def fake(profile, system, user, opts):
        calls.append(1)
        return ("not json" if len(calls) == 1 else '[{"statement": "The system shall export orders to CSV.", "type": "functional"}]'), {"in": 10, "out": 5}
    monkeypatch.setattr(pipeline.llm, "complete", fake)
    usage = {"in": 0, "out": 0}
    pipeline._ask({"kind": "openai"}, {"redact_pii": False}, "sys", "prompt", usage)
    assert usage["requests"] == 2 and usage["in"] == 20, "the retry for invalid JSON is a second request"


def _raw_story(title, feature, reqs, criteria=("Given a, when b, then c.",), epic="Orders"):
    return {"title": title, "epic": epic, "feature": feature, "narrative": "As a user, I want x, so that y.",
            "acceptance_criteria": list(criteria), "requirement_ids": list(reqs)}


def _req_list(*keys):
    return [{"key": k, "type": "functional", "statement": "The system shall do %s." % k, "epic": "Orders", "section": "S"} for k in keys]


def test_same_title_in_different_feature_areas_keeps_both_stories():
    out = pipeline._finalize([_raw_story("Manage users", "Billing", ["REQ-001"]), _raw_story("Manage users", "Returns", ["REQ-002"])], _req_list("REQ-001", "REQ-002"))
    assert [(s["title"], s["feature"], s["requirement_ids"]) for s in out] == [
        ("Manage users", "Billing", ["REQ-001"]), ("Manage users (Returns)", "Returns", ["REQ-002"])], "the second story is kept, with its links, under a title that tells them apart"


def test_same_title_in_the_same_feature_area_merges_so_no_requirement_loses_its_story():
    out = pipeline._finalize([_raw_story("Export report", "Reporting", ["REQ-001"], ["Given a, when b, then c."]),
                              _raw_story("Export Report", "Reporting", ["REQ-002", "REQ-001"], ["Given a, when b, then c.", "Given d, when e, then f."])],
                             _req_list("REQ-001", "REQ-002"))
    assert len(out) == 1 and out[0]["requirement_ids"] == ["REQ-001", "REQ-002"]
    assert out[0]["acceptance_criteria"] == ["Given a, when b, then c.", "Given d, when e, then f."], "criteria from the duplicate are kept, not repeated"


def test_unrelated_titles_and_invalid_links_are_unchanged():
    out = pipeline._finalize([_raw_story("Place an order", "Capture", ["REQ-001", "REQ-999"]), _raw_story("Cancel an order", "Capture", ["REQ-002"])], _req_list("REQ-001", "REQ-002"))
    assert [(s["title"], s["requirement_ids"]) for s in out] == [("Place an order", ["REQ-001"]), ("Cancel an order", ["REQ-002"])]


def test_a_requirement_stated_in_several_documents_keeps_every_source():
    docs = [{"id": "1", "filename": "v1.md", "text": "# Orders\n## Capture\nBR-101: The system shall allow buyers to place orders online.\nBR-102: The system shall email order confirmations to buyers."},
            {"id": "2", "filename": "addendum.docx", "text": "# Orders\n## Changes\nBR-201: The system shall allow buyers to place orders online.\nBR-202: The system shall block orders for suspended accounts."},
            {"id": "3", "filename": "legal.md", "text": "# Compliance\n## Audit\nBR-301: The system shall allow buyers to place orders online."}]
    opts = {"chunk_chars": 6000, "max_chunks_per_run": 50, "rag_top_k": 0, "redact_pii": True}
    reqs = pipeline.extract_requirements(docs, "generic", MOCK, opts, lambda m: None, {"in": 0, "out": 0})
    assert [r["source_id"] for r in reqs] == ["BR-101", "BR-102", "BR-202"], "still one requirement per statement"
    shared = reqs[0]["source_ref"]
    assert shared.startswith("v1.md > Capture") and "also addendum.docx > Changes (BR-201)" in shared and "also legal.md > Audit (BR-301)" in shared
    assert reqs[1]["source_ref"] == "v1.md > Capture", "a requirement stated once is unchanged"
    # two BRD ids in one document that state the same thing both stay traceable, and a repeat is never listed twice
    twice = [{"id": "1", "filename": "v1.md", "text": "# A\n## B\nBR-1: The system shall export orders.\nBR-2: The system shall export orders.\nBR-2: The system shall export orders."}]
    only = pipeline.extract_requirements(twice, "generic", MOCK, opts, lambda m: None, {"in": 0, "out": 0})
    assert len(only) == 1 and only[0]["source_ref"] == "v1.md > B; also v1.md > B (BR-2)"


# ---- near-duplicate detection on repetitive input ---------------------------------------------------------------

def _templated(n):
    return [{"key": "REQ-%d" % i, "statement": "The system shall allow an authorised %s user to process request number %d, validate it against the business rules "
             "for area %d, record who changed what and when, and notify the requester within %d seconds." % (["buyer", "approver", "auditor", "planner"][i % 4], i, i % 17, 5 + i % 55),
             "status": "draft"} for i in range(n)]


def test_duplicate_partners_are_capped_so_memory_does_not_grow_with_the_square_of_the_set():
    same = [{"key": "REQ-%d" % i, "statement": "The system shall export every invoice to CSV for the finance team.", "status": "draft"} for i in range(500)]
    dupes = quality.duplicate_pairs(same)
    assert len(dupes) == 500, "every one of the identical statements is flagged"
    assert max(len(v) for v in dupes.values()) <= quality.MAX_PARTNERS, "but none lists all 499 others (%d)" % max(len(v) for v in dupes.values())
    assert all(k not in v for k, v in dupes.items()), "an item is never its own duplicate"


def test_repetitive_requirements_do_not_make_duplicate_detection_slow():
    import time
    items = _templated(6000)
    t = time.time()
    dupes = quality.duplicate_pairs(items)
    took = time.time() - t
    assert took < 5, "6000 templated requirements took %.1f s (about a minute before the comparison budget)" % took
    assert isinstance(dupes, dict)


def test_the_comparison_budget_stops_the_work_and_what_was_found_so_far_is_still_right():
    items = _templated(1500) + [{"key": "DUP-A", "statement": "The system shall archive closed orders after seven years automatically.", "status": "draft"},
                                {"key": "DUP-B", "statement": "The system shall archive closed orders after seven years automatically.", "status": "draft"}]
    full = quality.duplicate_pairs(items, max_comparisons=10**9)
    tiny = quality.duplicate_pairs(items, max_comparisons=2000)
    assert "DUP-B" in full["DUP-A"], "the real pair is found with an ample budget"
    assert set(tiny) <= set(full) and all(set(v) <= set(full[k]) for k, v in tiny.items()), "a small budget finds fewer pairs, never wrong ones"


def test_small_sets_are_still_exact():
    items = [{"key": "A", "statement": "The system shall export invoices to CSV on request.", "status": "draft"},
             {"key": "B", "statement": "The system shall export invoices to CSV on request!", "status": "draft"},
             {"key": "C", "statement": "Customer service agents may place orders on behalf of clinics.", "status": "draft"}]
    assert quality.duplicate_pairs(items) == {"A": ["B"], "B": ["A"]}


def test_the_compliance_matrix_lists_each_requirement_under_each_tag_with_its_stories():
    import csv
    import io
    tagged = dict(REQ, key="REQ-001", compliance_tags=["SOX", "GDPR"])
    other = dict(REQ, key="REQ-002", compliance_tags=["GDPR"])
    plain = dict(REQ, key="REQ-003", compliance_tags=[])
    story = dict(_story(), requirement_ids=["REQ-001"])
    data, mime = exporters.compliance_csv([story], [tagged, other, plain])
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert rows[0][:3] == ["compliance_tag", "requirement", "type"] and mime.startswith("text/csv")
    assert [(r[0], r[1], r[-1]) for r in rows[1:]] == [("GDPR", "REQ-001", "covered"), ("GDPR", "REQ-002", "GAP"), ("SOX", "REQ-001", "covered")], "REQ-003 has no tag so it is not listed"


def test_the_excel_workbook_gets_a_compliance_sheet_only_when_tags_are_in_use():
    import io

    from openpyxl import load_workbook
    tagged = dict(REQ, key="REQ-001", compliance_tags=["SOX"])
    data, _ = exporters.xlsx([dict(_story(), requirement_ids=["REQ-001"])], [tagged, dict(REQ, key="REQ-002", compliance_tags=[])])
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Stories", "Requirements", "Traceability", "Compliance"]
    rows = list(wb["Compliance"].iter_rows(values_only=True))
    assert rows[0][:2] == ("compliance_tag", "requirement") and [(r[0], r[1], r[-1]) for r in rows[1:]] == [("SOX", "REQ-001", "covered")]
    plain, _ = exporters.xlsx([_story()], [dict(REQ, key="REQ-001", compliance_tags=[])])
    assert load_workbook(io.BytesIO(plain)).sheetnames == ["Stories", "Requirements", "Traceability"]


def test_the_html_report_summarises_compliance_coverage_only_when_tags_are_in_use():
    tagged = dict(REQ, key="REQ-001", compliance_tags=["SOX"])
    other = dict(REQ, key="REQ-002", compliance_tags=["SOX", "<b>GDPR</b>"])
    data, _ = exporters.html_report([dict(_story(), requirement_ids=["REQ-001"])], [tagged, other], "Demo")
    page = data.decode()
    assert "<h2>Compliance coverage</h2>" in page
    assert "<td>SOX</td><td>2</td><td>1 of 2" in page.replace("</td><td>", "</td><td>"), "SOX: two requirements, one covered"
    assert "<b>GDPR</b>" not in page and "&lt;b&gt;GDPR&lt;/b&gt;" in page, "tags are escaped"
    assert page.index("Compliance coverage") < page.index("Review and approval"), "it comes before the sign-off block"
    plain, _ = exporters.html_report([_story()], [dict(REQ, key="REQ-001", compliance_tags=[])], "Demo")
    assert "Compliance coverage" not in plain.decode()
