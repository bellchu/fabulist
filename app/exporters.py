"""Exporters. Each takes (stories, requirements) in canonical form and returns (bytes, mime)."""
import csv
import html
import io
import json
import re
import zipfile
from typing import Any, Callable, Dict, List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

Item = Dict[str, Any]
Output = Tuple[bytes, str]
STORY_COLUMNS = ["key", "epic", "feature", "title", "persona", "narrative", "acceptance_criteria", "priority",
                 "story_points", "requirement_ids", "labels", "product_scope", "variant_criteria", "compliance_tags",
                 "open_questions", "quality_score", "ready", "status"]
REQ_COLUMNS = ["key", "type", "statement", "epic", "section", "source_id", "source_ref", "actors", "product_scope",
               "compliance_tags", "quality_score", "status"]


def _cell(item: Item, col: str) -> str:
    v = item.get(col, "")
    if col == "variant_criteria":
        return "\n".join("[%s] %s" % (x.get("scope", ""), x.get("criteria", "")) for x in v or [])
    if isinstance(v, list):
        return "\n".join(str(x) for x in v)
    if isinstance(v, bool):
        return "yes" if v else "no"
    return "" if v is None else str(v)


def _guard(value: str) -> str:
    """Neutralize spreadsheet formula injection in CSV/XLSX cells."""
    return "'" + value if value[:1] in ("=", "+", "-", "@", "\t", "\r") else value


def _csv(header: List[str], rows: List[List[str]]) -> Output:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(header)
    for r in rows:
        w.writerow([_guard(c) for c in r])
    return buf.getvalue().encode("utf-8"), "text/csv"


def _rtm_rows(stories: List[Item], reqs: List[Item]) -> List[List[str]]:
    by_req: Dict[str, List[Item]] = {}
    for s in stories:
        for k in s.get("requirement_ids", []):
            by_req.setdefault(k, []).append(s)
    rows = []
    for r in reqs:
        linked = by_req.get(r["key"], [])
        rows.append([r["key"], r["type"], r["statement"], r.get("source_ref", ""),
                     "\n".join(s["key"] for s in linked), "\n".join(s["title"] for s in linked),
                     "covered" if linked else "GAP"])
    return rows


RTM_HEADER = ["requirement", "type", "statement", "source", "story_keys", "story_titles", "coverage"]


def stories_csv(stories: List[Item], reqs: List[Item]) -> Output:
    return _csv(STORY_COLUMNS, [[_cell(s, c) for c in STORY_COLUMNS] for s in stories])


def requirements_csv(stories: List[Item], reqs: List[Item]) -> Output:
    return _csv(REQ_COLUMNS, [[_cell(r, c) for c in REQ_COLUMNS] for r in reqs])


def rtm_csv(stories: List[Item], reqs: List[Item]) -> Output:
    return _csv(RTM_HEADER, _rtm_rows(stories, reqs))


def jira_csv(stories: List[Item], reqs: List[Item]) -> Output:
    rows = []
    for s in stories:
        desc = s["narrative"] + "\n\nh3. Acceptance criteria\n" + "\n".join("* " + a for a in s["acceptance_criteria"])
        if s.get("variant_criteria"):
            desc += "\n\nh3. Product variants\n" + "\n".join("* %s: %s" % (v["scope"], v["criteria"]) for v in s["variant_criteria"])
        if s.get("requirement_ids"):
            desc += "\n\nh3. Requirements\n" + ", ".join(s["requirement_ids"])
        labels = [x.replace(" ", "-") for x in s.get("labels", []) + s.get("compliance_tags", [])]
        rows.append([s["title"], "Story", s["priority"], str(s.get("story_points") or ""), desc, " ".join(labels), s["epic"]])
    return _csv(["Summary", "Issue Type", "Priority", "Story Points", "Description", "Labels", "Epic Link"], rows)


def azure_csv(stories: List[Item], reqs: List[Item]) -> Output:
    rows = []
    for s in stories:
        ac = "<br>".join(s["acceptance_criteria"])
        rows.append(["User Story", s["title"], s["narrative"], ac, {"High": "1", "Medium": "2", "Low": "3"}[s["priority"]],
                     str(s.get("story_points") or ""), "; ".join(s.get("labels", []) + s.get("compliance_tags", []))])
    return _csv(["Work Item Type", "Title", "Description", "Acceptance Criteria", "Priority", "Story Points", "Tags"], rows)


def as_json(stories: List[Item], reqs: List[Item]) -> Output:
    return json.dumps({"requirements": reqs, "stories": stories}, indent=2).encode("utf-8"), "application/json"


def markdown(stories: List[Item], reqs: List[Item]) -> Output:
    lines = ["# Product backlog", ""]
    epic = feature = None
    for s in stories:
        if s["epic"] != epic:
            epic, feature = s["epic"], None
            lines += ["## Epic: " + epic, ""]
        if s.get("feature") != feature:
            feature = s.get("feature")
            lines += ["### Feature: %s" % feature, ""]
        pts = " · %s pts" % s["story_points"] if s.get("story_points") else ""
        lines += ["#### %s: %s" % (s["key"], s["title"]), "",
                  "_%s · Priority %s%s · %s_" % (s["persona"], s["priority"], pts, s["status"]), "", s["narrative"], "",
                  "**Acceptance criteria**", ""]
        lines += ["- " + a for a in s["acceptance_criteria"]] or ["- _none_"]
        if s.get("variant_criteria"):
            lines += ["", "**Product variants**", ""] + ["- _%s_: %s" % (v["scope"], v["criteria"]) for v in s["variant_criteria"]]
        meta = []
        for key, label in (("requirement_ids", "Requirements"), ("product_scope", "Products"),
                           ("compliance_tags", "Compliance"), ("labels", "Labels")):
            if s.get(key):
                meta.append("%s: %s" % (label, ", ".join(s[key])))
        if meta:
            lines += ["", " | ".join(meta)]
        if s.get("open_questions"):
            lines += ["", "**Open questions**", ""] + ["- " + q for q in s["open_questions"]]
        lines.append("")
    return "\n".join(lines).encode("utf-8"), "text/markdown"


def xlsx(stories: List[Item], reqs: List[Item]) -> Output:
    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="334155")

    def sheet(ws: Any, header: List[str], rows: List[List[str]], widths: Dict[int, int]) -> None:
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
        for r in rows:
            ws.append([_guard(v) for v in r])
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        for idx, w in widths.items():
            ws.column_dimensions[chr(64 + idx)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "Stories"
    sheet(ws, STORY_COLUMNS, [[_cell(s, c) for c in STORY_COLUMNS] for s in stories], {4: 40, 6: 60, 7: 70})
    sheet(wb.create_sheet("Requirements"), REQ_COLUMNS, [[_cell(r, c) for c in REQ_COLUMNS] for r in reqs], {3: 80})
    sheet(wb.create_sheet("Traceability"), RTM_HEADER, _rtm_rows(stories, reqs), {3: 70, 6: 50})
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


_REPORT_CSS = """
body{font:14px/1.55 -apple-system,Segoe UI,Roboto,sans-serif;color:#111827;max-width:960px;margin:32px auto;padding:0 20px}
h1{font-size:26px;margin:0 0 4px}h2{font-size:19px;margin:36px 0 8px;padding-bottom:6px;border-bottom:2px solid #e5e7eb}
h3{font-size:15px;margin:22px 0 6px;color:#4338ca}h4{font-size:14px;margin:0 0 4px}.muted{color:#6b7280}
.kpis{display:grid;grid-template-columns:repeat(4,1fr);gap:10px;margin:18px 0}.kpi{border:1px solid #e5e7eb;border-radius:8px;padding:10px 12px}
.kpi b{display:block;font-size:22px}.story{border:1px solid #e5e7eb;border-radius:8px;padding:12px 14px;margin:10px 0;break-inside:avoid}
.tag{display:inline-block;font-size:11px;border:1px solid #d1d5db;border-radius:10px;padding:0 7px;margin:2px 3px 0 0;color:#374151}
.High{color:#b91c1c}.Medium{color:#b45309}.Low{color:#047857}ul{margin:6px 0 0 18px;padding:0}table{border-collapse:collapse;width:100%;font-size:13px}
td,th{border-bottom:1px solid #e5e7eb;padding:6px 8px;text-align:left;vertical-align:top}th{color:#6b7280;font-weight:600}.gap{color:#b91c1c;font-weight:600}
@media print{body{margin:0}h2{break-after:avoid}.story{page-break-inside:avoid}}
"""


def html_report(stories: List[Item], reqs: List[Item], title: str = "Product backlog") -> Output:
    e = html.escape
    covered = {k for s in stories for k in s.get("requirement_ids", [])}
    gaps = [r for r in reqs if r["key"] not in covered]
    points = sum(s.get("story_points") or 0 for s in stories)
    out = ["<!doctype html><html lang='en'><head><meta charset='utf-8'><title>%s</title><style>%s</style></head><body>" % (e(title), _REPORT_CSS),
           "<h1>%s</h1><div class='muted'>Generated by Fabulist</div>" % e(title),
           "<div class='kpis'><div class='kpi'><b>%d</b>stories</div><div class='kpi'><b>%d</b>requirements</div>"
           "<div class='kpi'><b>%d%%</b>coverage</div><div class='kpi'><b>%d</b>story points</div></div>"
           % (len(stories), len(reqs), round(100 * (len(reqs) - len(gaps)) / len(reqs)) if reqs else 0, points)]
    epic = feature = None
    for s in stories:
        if s["epic"] != epic:
            epic, feature = s["epic"], None
            out.append("<h2>Epic: %s</h2>" % e(epic))
        if s.get("feature") != feature:
            feature = s.get("feature")
            out.append("<h3>%s</h3>" % e(feature or ""))
        tags = "".join("<span class='tag'>%s</span>" % e(t) for t in s.get("requirement_ids", []) + s.get("compliance_tags", []) + s.get("product_scope", []))
        pts = " · %s pts" % s["story_points"] if s.get("story_points") else ""
        out.append("<div class='story'><h4>%s · %s</h4><div class='muted'><span class='%s'>%s</span> priority%s · %s</div><p>%s</p>"
                   % (e(s["key"]), e(s["title"]), e(s["priority"]), e(s["priority"]), pts, e(s["status"]), e(s["narrative"])))
        out.append("<b>Acceptance criteria</b><ul>%s</ul>" % "".join("<li>%s</li>" % e(a) for a in s["acceptance_criteria"]))
        if s.get("variant_criteria"):
            out.append("<b>Product variants</b><ul>%s</ul>" % "".join("<li><i>%s</i>: %s</li>" % (e(v["scope"]), e(v["criteria"])) for v in s["variant_criteria"]))
        if s.get("open_questions"):
            out.append("<b>Open questions</b><ul>%s</ul>" % "".join("<li>%s</li>" % e(q) for q in s["open_questions"]))
        out.append("<div>%s</div></div>" % tags)
    out.append("<h2>Traceability</h2><table><tr><th>Requirement</th><th>Statement</th><th>Stories</th></tr>")
    by_req: Dict[str, List[str]] = {}
    for s in stories:
        for k in s.get("requirement_ids", []):
            by_req.setdefault(k, []).append(s["key"])
    for r in reqs:
        linked = by_req.get(r["key"])
        out.append("<tr><td>%s</td><td>%s</td><td>%s</td></tr>" % (e(r["key"]), e(r["statement"]),
                                                                   e(", ".join(linked)) if linked else "<span class='gap'>Gap</span>"))
    out.append("</table></body></html>")
    return "".join(out).encode("utf-8"), "text/html"


_STEP = re.compile(r"\b(given|when|then|and|but)\b", re.I)


def gherkin_steps(criterion: str) -> List[str]:
    """Split 'Given x, when y, then z' (or a single 'Then z') into Gherkin steps."""
    text = criterion.strip().rstrip(".")
    starts = [m.start() for m in _STEP.finditer(text) if m.start() == 0 or text[m.start() - 1] in " ,;"]
    if not starts or starts[0] != 0:
        return ["Then " + text[0].lower() + text[1:]] if text else []
    parts = [text[a:b].strip(" ,;") for a, b in zip(starts, starts[1:] + [len(text)])]
    return [p[0].upper() + p[1:] for p in parts if p]


def _tag(value: str) -> str:
    return "@" + re.sub(r"[^\w.-]+", "-", value.strip()).strip("-")


def gherkin_zip(stories: List[Item], reqs: List[Item]) -> Output:
    """One .feature file per story feature; one scenario per acceptance criterion, tagged for traceability."""
    by_feature: Dict[str, List[Item]] = {}
    for s in stories:
        by_feature.setdefault("%s / %s" % (s.get("epic", "General"), s.get("feature", "General")), []).append(s)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, items in sorted(by_feature.items()):
            lines = ["Feature: %s" % name, ""]
            for s in items:
                tags = [_tag(s["key"])] + [_tag(r) for r in s.get("requirement_ids", [])] + [_tag(c) for c in s.get("compliance_tags", [])]
                lines += ["  # %s" % s.get("narrative", "")]
                for i, ac in enumerate(s.get("acceptance_criteria") or [], 1):
                    steps = gherkin_steps(ac)
                    if not steps:
                        continue
                    lines += ["  %s" % " ".join(tags), "  Scenario: %s: %s (%d)" % (s["key"], s["title"], i)]
                    lines += ["    " + st for st in steps] + [""]
                for v in s.get("variant_criteria") or []:
                    lines += ["  %s %s" % (" ".join(tags), _tag("variant")), "  Scenario: %s: %s [%s]" % (s["key"], s["title"], v["scope"]),
                              "    Given the product is in class \"%s\"" % v["scope"], "    Then " + v["criteria"].rstrip(".")[0].lower() + v["criteria"].rstrip(".")[1:], ""]
            fname = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-") or "feature"
            z.writestr("features/%s.feature" % fname, "\n".join(lines).rstrip() + "\n")
    return buf.getvalue(), "application/zip"


# key -> (function, file extension, label, description)
EXPORTERS: Dict[str, Tuple[Callable[[List[Item], List[Item]], Output], str, str, str]] = {
    "html": (html_report, "html", "Stakeholder report (HTML)", "Readable backlog with coverage summary; open in a browser and print to PDF."),
    "xlsx": (xlsx, "xlsx", "Excel workbook", "Stories, requirements and traceability matrix on separate sheets."),
    "jira-csv": (jira_csv, "csv", "Jira import (CSV)", "For Jira's external system CSV import, with an Epic Link column."),
    "azure-csv": (azure_csv, "csv", "Azure DevOps import (CSV)", "Work items for Azure Boards CSV import."),
    "md": (markdown, "md", "Markdown", "Readable backlog for Confluence, wikis or pull requests."),
    "gherkin": (gherkin_zip, "zip", "Test scenarios (Gherkin)", "Feature files from acceptance criteria, tagged with story, requirement and compliance IDs."),
    "stories-csv": (stories_csv, "csv", "Stories (CSV)", "Every story field, one row per story."),
    "requirements-csv": (requirements_csv, "csv", "Requirements (CSV)", "Extracted requirements with quality scores."),
    "rtm-csv": (rtm_csv, "csv", "Traceability matrix (CSV)", "Requirement-to-story coverage, with gaps flagged."),
    "json": (as_json, "json", "JSON", "Requirements and stories for integrations and APIs."),
}
