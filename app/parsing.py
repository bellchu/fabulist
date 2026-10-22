"""Extract plain text from uploaded BRD files."""
import io
from typing import Callable, Dict, List

from . import config


class ParseError(ValueError):
    pass


def _check_zip(data: bytes) -> None:
    """DOCX/XLSX are ZIP archives. The upload limit is on the compressed size, so check what they would
    expand to (from the archive index, without expanding anything) before handing them to a parser."""
    import zipfile

    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        infos = zf.infolist()
        if len(infos) > 10_000:
            raise ParseError("The file contains too many parts to be a normal Word or Excel document.")
        total = sum(i.file_size for i in infos)
        if total > config.MAX_UNZIPPED_MB * 1024 * 1024:
            raise ParseError("The file expands to %d MB, above the limit of %d MB. Split it or remove embedded content."
                             % (total // (1024 * 1024), config.MAX_UNZIPPED_MB))


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ParseError("The PDF is password-protected. Remove the password and upload it again.")
    if len(reader.pages) > config.MAX_PDF_PAGES:
        raise ParseError("The PDF has %d pages, above the limit of %d. Split it by section." % (len(reader.pages), config.MAX_PDF_PAGES))
    out, size = [], 0
    for page in reader.pages:
        text = page.extract_text() or ""
        out.append(text)
        size += len(text)
        if size > config.MAX_DOC_CHARS:  # the rest would be cut anyway; don't spend minutes extracting it
            break
    if not any(t.strip() for t in out):
        raise ParseError("No text could be extracted. This looks like a scanned PDF: upload a text-based export or run OCR first.")
    return "\n\n".join(out)


def _docx(data: bytes) -> str:
    from docx import Document

    _check_zip(data)
    doc = Document(io.BytesIO(data))
    # Style names by id, built once. `paragraph.style` looks the style up from scratch for every paragraph (it scans the style table):
    # 0.8 ms each, so a 6 MB document of 33,000 paragraphs spent 26 of its 27 seconds there, holding the interpreter for all other requests.
    names = {s.style_id: (s.name or "").lower() for s in doc.styles}
    out, size = [], 0
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = names.get(p._p.style, "")  # a paragraph with no style is Normal: not a heading
        out.append("\n# " + text if style.startswith("heading") or style == "title" else text)
        size += len(text)
        if size > config.MAX_DOC_CHARS:  # the rest would be cut anyway; don't spend minutes extracting it
            return "\n".join(out)
    for table in doc.tables:
        for row in table.rows[: MAX_ROWS + 1]:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                out.append(" | ".join(cells))
                size += sum(len(c) for c in cells)
        if size > config.MAX_DOC_CHARS:
            break
    return "\n".join(out)


_STATEMENT_COLS = ("requirement", "statement", "description", "user story", "story", "text", "detail")
_ID_COLS = ("id", "req id", "requirement id", "ref", "reference", "key", "#")
MAX_ROWS = 5000
MAX_SHEETS = 50


def _rows_to_text(title: str, rows: List[List[str]]) -> str:
    """Render a table as BRD text. If a requirement/description column exists, emit 'ID: statement'
    lines (plus other columns as context); otherwise 'header: value' pairs per row."""
    rows = [[c.strip() for c in r] for r in rows if any(c.strip() for c in r)]
    cut = len(rows) > MAX_ROWS + 1
    rows = rows[: MAX_ROWS + 1]
    if not rows:
        return ""
    header = [h.lower() for h in rows[0]]
    stmt = next((i for i, h in enumerate(header) if any(h == c or h.startswith(c) for c in _STATEMENT_COLS)), None)
    ident = next((i for i, h in enumerate(header) if h in _ID_COLS), None)
    out = ["# " + title]
    for r in rows[1:]:
        r = r + [""] * (len(header) - len(r))
        if stmt is not None and r[stmt]:
            extra = ["%s: %s" % (rows[0][i], v) for i, v in enumerate(r) if v and i not in (stmt, ident) and i < len(rows[0])]
            line = ("%s: " % r[ident] if ident is not None and r[ident] else "") + r[stmt].rstrip(".") + "."
            out.append(line + (" (" + "; ".join(extra) + ")" if extra else ""))
        elif stmt is None:
            out.append("; ".join("%s: %s" % (rows[0][i], v) for i, v in enumerate(r) if v and i < len(rows[0])))
    if cut:
        out.append("[Only the first %d rows of this sheet were read; split the file to extract the rest.]" % MAX_ROWS)
    return "\n".join(out)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    _check_zip(data)
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    if len(wb.worksheets) > MAX_SHEETS:
        parts.append("[Only the first %d of %d sheets were read.]" % (MAX_SHEETS, len(wb.worksheets)))
    for ws in wb.worksheets[:MAX_SHEETS]:
        rows = [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True, max_row=MAX_ROWS + 2)]
        parts.append(_rows_to_text(ws.title, rows))
    wb.close()
    return "\n\n".join(p for p in parts if p)


def _csv(data: bytes) -> str:
    import csv

    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return _rows_to_text("Requirements", list(csv.reader(io.StringIO(text), dialect)))


def _text(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


PARSERS: Dict[str, Callable[[bytes], str]] = {
    ".pdf": _pdf, ".docx": _docx, ".xlsx": _xlsx, ".csv": _csv, ".md": _text, ".markdown": _text, ".txt": _text,
}


def extract_text(filename: str, data: bytes) -> str:
    ext = "." + filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    parser = PARSERS.get(ext)
    if parser is None:
        raise ParseError("Unsupported file type. Use PDF, DOCX, XLSX, CSV, MD or TXT.")
    try:
        text = parser(data)
    except ParseError:
        raise
    except Exception as exc:  # corrupt files surface as a clean 4xx, not a 500
        if ext in (".docx", ".xlsx") or ext == ".pdf":
            raise ParseError("The file is damaged or is not a real %s file (%s)." % (ext[1:].upper(), type(exc).__name__)) from exc
        raise ParseError("Could not read file: %s" % type(exc).__name__) from exc
    text = text.replace("\x00", "").strip()
    if not text:
        raise ParseError("No text could be extracted from the file.")
    return clip(text)


def clip(text: str) -> str:
    """Cap a document at MAX_DOC_CHARS and say so in the text, so nobody assumes everything was read."""
    if len(text) <= config.MAX_DOC_CHARS:
        return text
    return text[: config.MAX_DOC_CHARS] + "\n\n[Truncated: this document is longer than %d characters; the rest was not read. Split it into parts.]" % config.MAX_DOC_CHARS
